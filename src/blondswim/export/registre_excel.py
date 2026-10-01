"""
Full de registre setmanal (Fase E+I): el nedador l'omple al llarg de la
setmana i després s'ingereix amb `ingestion.registre_setmana`.

Tres pestanyes:
- "Sessions": una fila per sessió planificada (natació i gimnàs) amb minuts
  reals, RPE (CR-10, 0-10), assoliment (1-5) i metres reals.
- "Benestar (SRSS)": un registre per dia d'entrenament, abans de la primera
  sessió (8 ítems de 0 a 6).
- "Sèrie de control": temps de cada 100, braçades per llargada i esforç.

L'sRPE es pot registrar fins i tot al vespre (és robust d'1 minut a 14 dies
després de la sessió).
"""

from datetime import timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from blondswim.agents.pla_setmanal import DIA_SERIE_CONTROL, ETIQUETA_ROL
from blondswim.models.franja import ETIQUETA_FRANJA, ETIQUETA_MODALITAT, ORDRE_FRANJA
from blondswim.models.macrocicle import Microcicle
from blondswim.models.nedador import Nedador
from blondswim.models.registre import (
    ETIQUETES_CR10,
    ITEMS_SRSS_ESTRES,
    ITEMS_SRSS_RECUPERACIO,
)
from blondswim.models.sessio import Sessio
from blondswim.utils.dates import parsejar_rang_dates

FULL_SESSIONS = "Sessions"
FULL_SRSS = "Benestar (SRSS)"
FULL_CONTROL = "Sèrie de control"

# Capçaleres (també les fa servir la ingestió; no canviar sense revisar-la).
COLUMNES_SESSIONS = [
    "Data", "Dia", "Franja", "Modalitat", "Sessió", "Planificat",
    "Minuts reals", "RPE (0-10)", "Assoliment (1-5)", "Metres reals", "Notes",
]
COLUMNES_SRSS = ["Data", "Dia", *ITEMS_SRSS_RECUPERACIO, *ITEMS_SRSS_ESTRES]
COLUMNES_CONTROL = [
    "Data", "100 #1", "100 #2", "100 #3", "100 #4",
    "Braçades/llargada", "RPE (0-10)", "Notes",
]
FILA_CAPCALERA = 3  # fila 1: títol, fila 2: nedador_id i instruccions

_DIES_ORDRE = {
    "dilluns": 0, "dimarts": 1, "dimecres": 2, "dijous": 3,
    "divendres": 4, "dissabte": 5, "diumenge": 6,
}


def _capcalera(ws, titol: str, nedador: Nedador, instruccions: str, columnes: list[str]):
    ws.cell(row=1, column=1, value=titol).font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value=nedador.id)
    ws.cell(row=2, column=2, value=instruccions).font = Font(italic=True)
    for i, nom in enumerate(columnes, start=1):
        ws.cell(row=FILA_CAPCALERA, column=i, value=nom).font = Font(bold=True)


def _validacio_enter(ws, minim: int, maxim: int, rang: str) -> None:
    dv = DataValidation(
        type="whole", operator="between", formula1=str(minim), formula2=str(maxim),
        allow_blank=True, showErrorMessage=True,
        error=f"Valor entre {minim} i {maxim}",
    )
    ws.add_data_validation(dv)
    dv.add(rang)


def _descripcio_sessio(sessio: Sessio) -> str:
    if sessio.modalitat != "natacio":
        return ETIQUETA_MODALITAT[sessio.modalitat]
    return ETIQUETA_ROL.get(sessio.rol or "", sessio.tipus_sessio)


def _planificat(sessio: Sessio) -> str:
    if sessio.modalitat == "natacio":
        return f"{sessio.volum_total} m"
    return f"{sessio.durada_min} min" if sessio.durada_min else ""


def exportar_registre_setmana(
    nedador: Nedador,
    microcicle: Microcicle,
    sessions: list[Sessio],
    output_path: Path,
) -> Path:
    """Crea el full de registre en blanc d'una setmana."""
    dilluns, _ = parsejar_rang_dates(microcicle.dates)
    ordenades = sorted(
        (s for s in sessions if s.dia in _DIES_ORDRE),
        key=lambda s: (_DIES_ORDRE[s.dia], ORDRE_FRANJA[s.franja]),
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()

    # --- Sessions ---
    ws = wb.active
    ws.title = FULL_SESSIONS
    _capcalera(
        ws,
        f"Registre setmana {microcicle.setmana} ({microcicle.dates}) — {nedador.nom}",
        nedador,
        "Omple minuts i RPE de cada sessió feta (pots fer-ho al vespre). "
        "Deixa la fila buida si no l'has feta.",
        COLUMNES_SESSIONS,
    )
    fila = FILA_CAPCALERA + 1
    for s in ordenades:
        data = dilluns + timedelta(days=_DIES_ORDRE[s.dia])
        valors = [
            data.isoformat(), s.dia, ETIQUETA_FRANJA[s.franja],
            ETIQUETA_MODALITAT[s.modalitat], _descripcio_sessio(s), _planificat(s),
        ]
        for col, valor in enumerate(valors, start=1):
            ws.cell(row=fila, column=col, value=valor)
        fila += 1
    darrera = max(fila - 1, FILA_CAPCALERA + 1)
    _validacio_enter(ws, 0, 10, f"H{FILA_CAPCALERA + 1}:H{darrera}")
    _validacio_enter(ws, 1, 5, f"I{FILA_CAPCALERA + 1}:I{darrera}")

    fila += 1
    ws.cell(row=fila, column=1, value="Escala RPE (CR-10)").font = Font(bold=True)
    for valor, etiqueta in ETIQUETES_CR10.items():
        fila += 1
        ws.cell(row=fila, column=1, value=valor)
        ws.cell(row=fila, column=2, value=etiqueta)

    # --- SRSS ---
    ws = wb.create_sheet(FULL_SRSS)
    _capcalera(
        ws,
        "Benestar abans de la primera sessió del dia",
        nedador,
        "0 = gens, 6 = totalment. Com et trobes ara?",
        COLUMNES_SRSS,
    )
    dies_entrenament = sorted({s.dia for s in ordenades}, key=_DIES_ORDRE.__getitem__)
    for i, dia in enumerate(dies_entrenament):
        data = dilluns + timedelta(days=_DIES_ORDRE[dia])
        ws.cell(row=FILA_CAPCALERA + 1 + i, column=1, value=data.isoformat())
        ws.cell(row=FILA_CAPCALERA + 1 + i, column=2, value=dia)
    if dies_entrenament:
        _validacio_enter(
            ws, 0, 6, f"C{FILA_CAPCALERA + 1}:J{FILA_CAPCALERA + len(dies_entrenament)}"
        )

    # --- Sèrie de control ---
    ws = wb.create_sheet(FULL_CONTROL)
    _capcalera(
        ws,
        "Sèrie de control 4x100 crol A2",
        nedador,
        "Temps de cada 100 com a text (1:40.5 o 100.5), braçades per llargada "
        "(mitjana) i esforç (0-10).",
        COLUMNES_CONTROL,
    )
    if any(s.dia == DIA_SERIE_CONTROL and s.modalitat == "natacio" for s in ordenades):
        data = dilluns + timedelta(days=_DIES_ORDRE[DIA_SERIE_CONTROL])
        ws.cell(row=FILA_CAPCALERA + 1, column=1, value=data.isoformat())
    for col in "BCDE":
        for row in range(FILA_CAPCALERA + 1, FILA_CAPCALERA + 3):
            ws[f"{col}{row}"].number_format = "@"  # text: Excel no ho converteix a hora
    _validacio_enter(ws, 0, 10, f"G{FILA_CAPCALERA + 1}:G{FILA_CAPCALERA + 2}")

    for full in wb.worksheets:
        for col in full.columns:
            amplada = max((len(str(c.value)) for c in col[2:] if c.value), default=8)
            full.column_dimensions[col[0].column_letter].width = min(amplada + 2, 30)

    wb.save(output_path)
    return output_path
