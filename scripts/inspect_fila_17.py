"""Script per inspeccionar les files al voltant de la fila 17."""

from pathlib import Path
import openpyxl

def llegir_capçaleres(ws, header_row: int = 1) -> dict[str, int]:
    """Llegir capçaleres (còpia de la funció original)."""
    capçaleres = {}
    for cell in ws[header_row]:
        if cell.value:
            # Normalitzar: minúscules, sense espais interns
            nom_normalitzat = str(cell.value).strip().lower().replace(" ", "")
            capçaleres[nom_normalitzat] = cell.column
    return capçaleres

def main():
    base_dir = Path(__file__).parent.parent
    fitxer = base_dir / "data" / "raw" / "Planificacio_Mesocicles_Jep.xlsx"
    
    print(f"📂 Obrint: {fitxer}")
    wb = openpyxl.load_workbook(fitxer, data_only=True)
    ws = wb["Microcicles"]
    
    # Llegir capçaleres (fila 2)
    capçaleres = llegir_capçaleres(ws, header_row=2)
    
    # Llegir valor de Meso a la fila 17
    meso_fila_17 = ws.cell(row=17, column=capçaleres["meso"]).value
    print(f"\n🔍 Fila 17 té Meso = '{meso_fila_17}'")
    
    # Trobar totes les files amb el mateix Meso
    print(f"\n📋 Files amb Meso = '{meso_fila_17}':")
    print(f"{'Fila':<6} {'Setmana':<10} {'Meso':<8} {'Tipus de setmana':<40} {'Volum':<10}")
    print("-" * 80)
    
    for num_fila in range(3, 50):  # Buscar fins a la fila 50
        setmana = ws.cell(row=num_fila, column=capçaleres["setmana"]).value
        if not setmana:
            break  # Acabar quan no hi ha més dades
        
        meso = ws.cell(row=num_fila, column=capçaleres["meso"]).value
        tipus_setmana = ws.cell(row=num_fila, column=capçaleres["tipusdesetmana"]).value
        volum = ws.cell(row=num_fila, column=capçaleres["volumobjectiu(m)"]).value
        
        # Mostrar files del mateix mesocicle
        if str(meso).strip() == str(meso_fila_17).strip():
            marca = ">>> " if num_fila == 17 else "    "
            print(f"{marca}{num_fila:<6} {str(setmana):<10} {str(meso):<8} {str(tipus_setmana):<40} {str(volum):<10}")

if __name__ == "__main__":
    main()
