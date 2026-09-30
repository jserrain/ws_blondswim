"""Script per debugar les capçaleres de la pestanya Microcicles."""

from pathlib import Path

import openpyxl

def llegir_capçaleres(ws, header_row: int = 1) -> dict[str, int]:
    """Llegir capçaleres (còpia de la funció original)."""
    capçaleres = {}
    for cell in ws[header_row]:
        if cell.value:
            # Normalitzar: minúscules, sense espais interns
            nom_normalitzat = str(cell.value).strip().lower().replace(" ", "")
            capçaleres[nom_normalitzat] = cell.column
    return capçaleres

def main():
    base_dir = Path(__file__).parent.parent
    fitxer = base_dir / "data" / "raw" / "Planificacio_Mesocicles_Jep.xlsx"
    
    print(f"📂 Obrint: {fitxer}")
    wb = openpyxl.load_workbook(fitxer, data_only=True)
    
    if "Microcicles" not in wb.sheetnames:
        print("❌ Pestanya 'Microcicles' no trobada")
        print(f"Pestanyes disponibles: {wb.sheetnames}")
        return
    
    ws = wb["Microcicles"]
    print("\n✅ Pestanya 'Microcicles' trobada")
    
    # Llegir capçaleres
    capçaleres = llegir_capçaleres(ws, header_row=1)
    
    print("\n📋 Capçaleres normalitzades retornades per llegir_capçaleres():")
    print(f"   Total: {len(capçaleres)} capçaleres")
    for nom, idx in sorted(capçaleres.items()):
        print(f"   '{nom}' -> columna {idx}")
    
    print("\n🔍 Contingut literal de les files 1, 2 i 3:")
    for fila_num in [1, 2, 3]:
        print(f"\n   Fila {fila_num}:")
        fila_buida = True
        for cell in ws[fila_num]:
            if cell.value:
                fila_buida = False
                print(f"      Columna {cell.column} ({cell.column_letter}): '{cell.value}'")
        if fila_buida:
            print("      (fila buida)")
    
    print("\n🔎 Verificació de capçaleres requerides:")
    required = ["mesocicle", "microcicle", "setmana", "dates", "tipus", "volumobjectiu(m)"]
    for req in required:
        if req in capçaleres:
            print(f"   ✓ '{req}' trobada")
        else:
            print(f"   ✗ '{req}' NO trobada")
            # Buscar similars
            similars = [k for k in capçaleres if req[:5] in k or k[:5] in req]
            if similars:
                print(f"      Similars: {similars}")

def debug_macrocicle():
    """Debugar la pestanya Macrocicle."""
    base_dir = Path(__file__).parent.parent
    fitxer = base_dir / "data" / "raw" / "Planificacio_Mesocicles_Jep.xlsx"
    
    print(f"\n{'='*60}")
    print("PESTANYA MACROCICLE")
    print(f"{'='*60}")
    
    wb = openpyxl.load_workbook(fitxer, data_only=True)
    
    if "Macrocicle" not in wb.sheetnames:
        print("❌ Pestanya 'Macrocicle' no trobada")
        return
    
    ws = wb["Macrocicle"]
    print("\n✅ Pestanya 'Macrocicle' trobada")
    
    print("\n🔍 Contingut literal de les files 1, 2 i 3:")
    for fila_num in [1, 2, 3]:
        print(f"\n   Fila {fila_num}:")
        fila_buida = True
        for cell in ws[fila_num]:
            if cell.value:
                fila_buida = False
                print(f"      Columna {cell.column} ({cell.column_letter}): '{cell.value}'")
        if fila_buida:
            print("      (fila buida)")
    
    # Buscar capçaleres a les files 4, 5, 6
    print("\n🔍 Contingut literal de les files 4, 5 i 6 (possibles capçaleres):")
    for fila_num in [4, 5, 6]:
        print(f"\n   Fila {fila_num}:")
        fila_buida = True
        for cell in ws[fila_num]:
            if cell.value:
                fila_buida = False
                print(f"      Columna {cell.column} ({cell.column_letter}): '{cell.value}'")
        if fila_buida:
            print("      (fila buida)")
    
    # Llegir capçaleres de la fila 6 (segons el codi actual)
    capçaleres = llegir_capçaleres(ws, header_row=6)
    print("\n📋 Capçaleres normalitzades de la fila 6:")
    print(f"   Total: {len(capçaleres)} capçaleres")
    for nom, idx in sorted(capçaleres.items()):
        print(f"   '{nom}' -> columna {idx}")
    
    # Mostrar primeres 2 files de dades (files 7 i 8)
    print("\n📄 Primeres 2 files de dades (files 7 i 8):")
    for fila_num in [7, 8]:
        print(f"\n   Fila {fila_num}:")
        fila_buida = True
        for cell in ws[fila_num]:
            if cell.value:
                fila_buida = False
                valor_str = str(cell.value)
                if len(valor_str) > 50:
                    valor_str = valor_str[:47] + "..."
                print(f"      Columna {cell.column} ({cell.column_letter}): '{valor_str}'")
        if fila_buida:
            print("      (fila buida)")
            break


if __name__ == "__main__":
    main()
    debug_macrocicle()
