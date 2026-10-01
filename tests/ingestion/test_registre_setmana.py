"""Tests del full de registre setmanal (exportació i ingestió, Fase E+I)."""

from datetime import date, time

import openpyxl
import pytest

from blondswim.agents.esquelet_sessions import generar_esquelet_sessions
from blondswim.export.registre_excel import (
    FILA_CAPCALERA,
    FULL_CONTROL,
    FULL_SESSIONS,
    FULL_SRSS,
    exportar_registre_setmana,
)
from blondswim.ingestion.registre_setmana import (
    carregar_registres,
    convertir_registre_setmana,
    parsejar_temps_100,
)
from blondswim.models.franja import SlotSessio
from blondswim.models.macrocicle import Microcicle
from blondswim.models.nedador import Nedador, RitmesCSS

DIES_5 = ["dilluns", "dimarts", "dimecres", "dijous", "divendres"]


@pytest.fixture
def nedador() -> Nedador:
    tipus = {dia: [SlotSessio(franja="mati", modalitat="natacio")] for dia in DIES_5}
    tipus["dilluns"].append(SlotSessio(franja="tarda", modalitat="gimnas", durada_min=60))
    return Nedador(
        id="jep",
        nom="Jep",
        categoria="master",
        proves_objectiu=["100m lliure"],
        mode_ritme="temps",
        ritmes_css=RitmesCSS(font="css_test", a2=82.0),
        setmana_tipus=tipus,
    )


@pytest.fixture
def microcicle() -> Microcicle:
    return Microcicle(
        setmana=41,
        dates="05-11/10/2026",
        mesocicle_id="meso_1",
        tipus_base="carrega",
        volum_objectiu=13600,
        dies_qualitat=False,
        test_css=False,
    )


@pytest.fixture
def full(nedador, microcicle, tmp_path):
    sessions = generar_esquelet_sessions(nedador, microcicle)
    return exportar_registre_setmana(
        nedador, microcicle, sessions, tmp_path / "registre_jep_2026-W41.xlsx"
    )


def _fila_sessio(ws, data: str, modalitat: str) -> int:
    for fila in range(FILA_CAPCALERA + 1, ws.max_row + 1):
        if ws.cell(fila, 1).value == data and ws.cell(fila, 4).value == modalitat:
            return fila
    raise AssertionError(f"Fila {data} {modalitat} no trobada")


def test_full_en_blanc(full):
    wb = openpyxl.load_workbook(full)
    ws = wb[FULL_SESSIONS]
    files = [
        [ws.cell(f, c).value for c in range(1, 7)]
        for f in range(FILA_CAPCALERA + 1, FILA_CAPCALERA + 7)
    ]
    assert files[0][:5] == ["2026-10-05", "dilluns", "matí", "Natació", "Aeròbic i tècnica"]
    assert files[1] == ["2026-10-05", "dilluns", "tarda", "Gimnàs", "Gimnàs", "60 min"]
    assert files[2][4] == "Qualitat"
    # SRSS: un dia per dia d'entrenament; sèrie de control el dimecres.
    srss = wb[FULL_SRSS]
    assert [srss.cell(f, 1).value for f in range(FILA_CAPCALERA + 1, FILA_CAPCALERA + 6)] == [
        "2026-10-05", "2026-10-06", "2026-10-07", "2026-10-08", "2026-10-09",
    ]
    assert wb[FULL_CONTROL].cell(FILA_CAPCALERA + 1, 1).value == "2026-10-07"
    # El full en blanc no té cap dada.
    buit = convertir_registre_setmana(full)
    assert (buit.sessions, buit.srss, buit.controls) == ([], [], [])


def test_ingestio_full_omplert(full):
    wb = openpyxl.load_workbook(full)
    ws = wb[FULL_SESSIONS]
    f = _fila_sessio(ws, "2026-10-05", "Natació")
    ws.cell(f, 7, 70)
    ws.cell(f, 8, 4)
    ws.cell(f, 9, 4)
    ws.cell(f, 10, 2900)
    f = _fila_sessio(ws, "2026-10-05", "Gimnàs")
    ws.cell(f, 7, 55)
    ws.cell(f, 8, 6)
    f = _fila_sessio(ws, "2026-10-06", "Natació")
    ws.cell(f, 7, 65)  # sense RPE: compta com a incompleta

    srss = wb[FULL_SRSS]
    for col, valor in enumerate([5, 4, 5, 4, 1, 2, 1, 1], start=3):
        srss.cell(FILA_CAPCALERA + 1, col, valor)

    control = wb[FULL_CONTROL]
    for col, valor in enumerate(["1:40.5", "1'40", 99.5, time(1, 39)], start=2):
        control.cell(FILA_CAPCALERA + 1, col, valor)
    control.cell(FILA_CAPCALERA + 1, 6, 18.5)
    control.cell(FILA_CAPCALERA + 1, 7, 4)
    wb.save(full)

    registre = convertir_registre_setmana(full)

    assert [(s.data, s.modalitat, s.franja, s.carrega) for s in registre.sessions] == [
        ("2026-10-05", "natacio", "mati", 280),
        ("2026-10-05", "gimnas", "tarda", 330),
        ("2026-10-06", "natacio", "mati", None),
    ]
    assert registre.sessions[0].volum_total_m == 2900
    assert registre.sessions[0].assoliment == 4
    assert len(registre.srss) == 1
    assert registre.srss[0].data == date(2026, 10, 5)
    assert registre.srss[0].mitjana_recuperacio == 4.5
    assert registre.controls[0].temps_100 == [100.5, 100.0, 99.5, 99.0]
    assert registre.controls[0].bracades_mitja == 18.5
    assert registre.controls[0].nedador_id == "jep"


def test_rpe_sense_minuts_es_error(full):
    wb = openpyxl.load_workbook(full)
    ws = wb[FULL_SESSIONS]
    ws.cell(_fila_sessio(ws, "2026-10-05", "Natació"), 8, 5)
    wb.save(full)
    with pytest.raises(ValueError, match="falten els minuts"):
        convertir_registre_setmana(full)


def test_srss_a_mitges_es_error(full):
    wb = openpyxl.load_workbook(full)
    wb[FULL_SRSS].cell(FILA_CAPCALERA + 1, 3, 5)
    wb.save(full)
    with pytest.raises(ValueError, match="8 ítems"):
        convertir_registre_setmana(full)


def test_carregar_registres_directori(full, tmp_path):
    wb = openpyxl.load_workbook(full)
    ws = wb[FULL_SESSIONS]
    f = _fila_sessio(ws, "2026-10-05", "Natació")
    ws.cell(f, 7, 70)
    ws.cell(f, 8, 4)
    wb.save(full)
    total = carregar_registres(tmp_path)
    assert len(total.sessions) == 1


@pytest.mark.parametrize(
    ("valor", "segons"),
    [(100.5, 100.5), ("1:40.5", 100.5), ("1'40", 100.0), ("99,5", 99.5), (time(1, 40), 100.0)],
)
def test_parsejar_temps_100(valor, segons):
    assert parsejar_temps_100(valor) == pytest.approx(segons)


def test_parsejar_temps_invalid():
    with pytest.raises(ValueError, match="no reconegut"):
        parsejar_temps_100("ràpid")
