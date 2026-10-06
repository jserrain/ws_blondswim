"""Pestanya «Sèries objectiu» del full de registre: exportació i lectura."""

from datetime import date

import openpyxl

from blondswim.agents import ritme_objectiu as ro
from blondswim.export.registre_excel import (
    FILA_CAPCALERA,
    FULL_OBJECTIU,
    afegir_full_objectiu,
    exportar_registre_setmana,
)
from blondswim.ingestion.registre_setmana import convertir_registre_setmana
from blondswim.models.macrocicle import Microcicle
from blondswim.models.nedador import Nedador


def _presc() -> ro.Prescripcio:
    return ro.Prescripcio(prova="100m Lliure", dia="dimarts", series=6, distancia=50,
                          descans="d/1:00", objectiu=40.0, execucio="x", pas=0, passos=8,
                          resum="r")


def _setup(tmp_path):
    n = Nedador(id="jep", nom="Jep", categoria="master", mode_ritme="temps",
                proves_objectiu=["100m Lliure"])
    micro = Microcicle(setmana=41, dates="05-11/10/2026", mesocicle_id="m",
                       tipus_base="carrega", volum_objectiu=13600, dies_qualitat=False,
                       test_css=False)
    return n, micro


def test_exportar_omplir_i_llegir(tmp_path):
    n, micro = _setup(tmp_path)
    path = exportar_registre_setmana(n, micro, [], tmp_path / "r.xlsx", [_presc()])
    wb = openpyxl.load_workbook(path)
    ws = wb[FULL_OBJECTIU]
    fila = FILA_CAPCALERA + 1
    assert ws.cell(row=fila, column=1).value == "2026-10-06"
    assert ws.cell(row=fila, column=4).value == 40.0
    for col, temps in enumerate(["39.8", "0:40.1", "40,0"], start=5):
        ws.cell(row=fila, column=col, value=temps)
    wb.save(path)
    [reg] = convertir_registre_setmana(path).objectius
    assert reg.data == date(2026, 10, 6) and reg.prova == "100m Lliure"
    assert reg.objectiu == 40.0 and reg.temps == [39.8, 40.1, 40.0]


def test_afegir_la_pestanya_a_un_full_antic(tmp_path):
    n, micro = _setup(tmp_path)
    path = exportar_registre_setmana(n, micro, [], tmp_path / "r.xlsx")
    wb = openpyxl.load_workbook(path)
    del wb[FULL_OBJECTIU]
    wb.save(path)
    assert convertir_registre_setmana(path).objectius == []
    assert afegir_full_objectiu(path, n, micro, [_presc()])
    assert not afegir_full_objectiu(path, n, micro, [_presc()])
    assert FULL_OBJECTIU in openpyxl.load_workbook(path).sheetnames
