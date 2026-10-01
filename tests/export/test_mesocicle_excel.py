"""Tests per a l'exportació de mesocicles a Excel (format per exercici)."""

import openpyxl
import pytest

from blondswim.export.mesocicle_excel import exportar_mesocicle_excel, exportar_setmana_excel
from blondswim.models.macrocicle import Mesocicle, Microcicle
from blondswim.models.nedador import Nedador, RitmesCSS
from blondswim.models.sessio import EstructuraSessio, Exercici, PartSessio, Sessio


@pytest.fixture
def nedador_test() -> Nedador:
    """Nedador de test."""
    return Nedador(
        id="test",
        nom="Test Nedador",
        edat=25,
        categoria="absolut",
        proves_objectiu=["200m lliure"],
        mode_ritme="temps",
        ritmes_css=RitmesCSS(
            font="css_test",
            data_test="2026-09-01",
            recuperacio=90.0,
            a1=85.0,
            a2=80.0,
            a3=75.0,
            velocitat=65.0,
        ),
    )


@pytest.fixture
def mesocicle_test() -> Mesocicle:
    """
    Mesocicle de test amb 2 setmanes:
    - Setmana 1: 28/09-04/10/2026 (creua de mes, ISO 40)
    - Setmana 2: 05-11/10/2026 (mateix mes, ISO 41)
    """
    microcicle_1 = Microcicle(
        setmana=1,
        dates="28/09-04/10/2026",
        mesocicle_id="meso1",
        tipus_base="carrega",
        volum_objectiu=15000,
        dies_qualitat=False,
        test_css=False,
    )
    microcicle_2 = Microcicle(
        setmana=2,
        dates="05-11/10/2026",
        mesocicle_id="meso1",
        tipus_base="qualitat",
        volum_objectiu=14000,
        dies_qualitat=True,
        test_css=False,
    )
    return Mesocicle(
        id="meso1",
        nom="Mesocicle 1",
        setmanes="1-2",
        dates="28/09-11/10/2026",
        tipus="Build1",
        fase_objectiu="Build1",
        metodologia_dominant="Polaritzat",
        volum_min=13000,
        volum_max=16000,
        volum_mitja_previst=14500,
        microcicles=[microcicle_1, microcicle_2],
    )


@pytest.fixture
def estructura_test() -> EstructuraSessio:
    """Estructura de sessió de test amb 1 part i 2 exercicis."""
    return EstructuraSessio(
        parts=[
            PartSessio(
                nom="Escalfament",
                percentatge_carrega=10,
                percentatge_qualitat=10,
                percentatge_descarrega=15,
                exercicis=[
                    Exercici(
                        series=1,
                        distancia_m=200,
                        execucio="N suau",
                        intensitat="Recuperació",
                        objectiu="Escalfament",
                    ),
                    Exercici(
                        series=4,
                        distancia_m=50,
                        execucio="Crol tècnica",
                        descans="c/15\"",
                        material="Pales petites",
                        intensitat="A1",
                        objectiu="Tècnica captura",
                    ),
                ],
            ),
        ]
    )


def test_exportar_mesocicle_capçalera_setmana_creua_mes(
    nedador_test, mesocicle_test, estructura_test, tmp_path
):
    """Setmana 1 (28/09-04/10) ha de mostrar Setmana ISO 40 i 'Setembre/Octubre'."""
    sessio = Sessio(
        id="test_1_dilluns",
        microcicle_setmana=1,
        dia="dilluns",
        tipus_sessio="carrega",
        volum_total=400,
        es_dia_opcional=False,
        estructura=estructura_test,
    )

    output_path = tmp_path / "mesocicle_test.xlsx"
    fitxer_creat = exportar_mesocicle_excel(
        nedador=nedador_test,
        mesocicle=mesocicle_test,
        resultats={1: [sessio]},
        output_path=output_path,
    )

    assert fitxer_creat.exists()

    wb = openpyxl.load_workbook(fitxer_creat)
    ws = wb.active

    capçalera_setmana = ws.cell(row=1, column=1).value
    assert "Setmana 40" in capçalera_setmana
    assert "Setembre/Octubre" in capçalera_setmana
    assert "Mesocicle 1" in capçalera_setmana
    assert "Build1" in capçalera_setmana

    assert ws.cell(row=1, column=1).font.bold is True

    # Capçalera de dia: "Dilluns 28" (sense mes)
    assert ws.cell(row=2, column=1).value == "Dilluns 28"
    assert ws.cell(row=2, column=1).font.bold is True

    # Capçaleres de columna
    assert ws.cell(row=3, column=2).value == "Treball"
    assert ws.cell(row=3, column=9).value == "Volum (m)"

    # Primer exercici (series=1 -> Treball = distancia sola)
    assert ws.cell(row=4, column=2).value == "200"
    assert ws.cell(row=4, column=3).value == "N suau"
    assert ws.cell(row=4, column=6).value == "Recuperació"
    assert ws.cell(row=4, column=9).value == 200

    # Segon exercici (series=4 -> Treball = "4x50")
    assert ws.cell(row=5, column=2).value == "4x50"
    assert ws.cell(row=5, column=4).value == "c/15\""
    assert ws.cell(row=5, column=5).value == "Pales petites"
    assert ws.cell(row=5, column=9).value == 200

    # Fila Total: volum sumat, múltiple de 25
    assert ws.cell(row=6, column=2).value == "Total"
    assert ws.cell(row=6, column=9).value == 400
    assert ws.cell(row=6, column=9).value % 25 == 0


def test_exportar_mesocicle_setmana_mateix_mes(
    nedador_test, mesocicle_test, estructura_test, tmp_path
):
    """Setmana 2 (05-11/10) ha de mostrar Setmana ISO 41 i només 'Octubre'."""
    sessio = Sessio(
        id="test_2_dimarts",
        microcicle_setmana=2,
        dia="dimarts",
        tipus_sessio="qualitat",
        volum_total=400,
        es_dia_opcional=False,
        estructura=estructura_test,
    )

    output_path = tmp_path / "mesocicle_test2.xlsx"
    fitxer_creat = exportar_mesocicle_excel(
        nedador=nedador_test,
        mesocicle=mesocicle_test,
        resultats={2: [sessio]},
        output_path=output_path,
    )

    wb = openpyxl.load_workbook(fitxer_creat)
    ws = wb.active

    capçalera_setmana = ws.cell(row=1, column=1).value
    assert "Setmana 41" in capçalera_setmana
    assert "Octubre" in capçalera_setmana
    assert "Setembre" not in capçalera_setmana

    # Dimarts = offset 1 des de dilluns 05/10 -> 6
    assert ws.cell(row=2, column=1).value == "Dimarts 6"


def test_exportar_mesocicle_omet_dies_abans_de_sessions_des_de(
    nedador_test, estructura_test, tmp_path
):
    """El 01/10/2026 surt com 'Dijous 1' i no hi apareixen els dies 28, 29 ni 30/09."""
    from datetime import date

    microcicle = Microcicle(
        setmana=40,
        dates="28/09-04/10/2026",
        mesocicle_id="meso1",
        tipus_base="carrega",
        volum_objectiu=12000,
        dies_qualitat=False,
        test_css=True,
        sessions_des_de=date(2026, 10, 1),
    )
    mesocicle = Mesocicle(
        id="meso1",
        nom="Base 1",
        setmanes="40-42",
        dates="28/09/2026-18/10/2026",
        tipus="Base",
        fase_objectiu="Base",
        metodologia_dominant="",
        volum_min=12000,
        volum_max=15000,
        volum_mitja_previst=13500,
        microcicles=[microcicle],
    )

    def _sessio(dia: str) -> Sessio:
        return Sessio(
            id=f"meso1_s40_{dia}",
            microcicle_setmana=40,
            dia=dia,
            tipus_sessio="carrega",
            volum_total=400,
            es_dia_opcional=False,
            estructura=estructura_test,
        )

    resultats = {
        40: [
            _sessio("dilluns"),
            _sessio("dimarts"),
            _sessio("dimecres"),
            _sessio("dijous"),
        ]
    }

    output_path = tmp_path / "mesocicle_sessions_des_de.xlsx"
    exportar_mesocicle_excel(
        nedador=nedador_test,
        mesocicle=mesocicle,
        resultats=resultats,
        output_path=output_path,
    )

    wb = openpyxl.load_workbook(output_path)
    ws = wb.active
    valors = [str(c.value) for row in ws.iter_rows() for c in row if c.value]

    assert any("Dijous 1" in v for v in valors)
    assert not any("Dilluns 28" in v for v in valors)
    assert not any("Dimarts 29" in v for v in valors)
    assert not any("Dimecres 30" in v for v in valors)


def test_exportar_mesocicle_resultats_buit(nedador_test, mesocicle_test, tmp_path):
    """Amb resultats buit no es genera cap fila."""
    output_path = tmp_path / "mesocicle_buit.xlsx"
    fitxer_creat = exportar_mesocicle_excel(
        nedador=nedador_test,
        mesocicle=mesocicle_test,
        resultats={},
        output_path=output_path,
    )

    assert fitxer_creat.exists()

    wb = openpyxl.load_workbook(fitxer_creat)
    ws = wb.active

    assert ws.max_row == 1
    assert ws.cell(row=1, column=1).value is None


# --- G6: exportació d'una sola setmana ---


def _sessio(setmana: int, dia: str, estructura: EstructuraSessio) -> Sessio:
    return Sessio(
        id=f"test_{setmana}_{dia}",
        microcicle_setmana=setmana,
        dia=dia,
        tipus_sessio="qualitat",
        volum_total=400,
        es_dia_opcional=False,
        estructura=estructura,
    )


def _valors(path) -> list[tuple]:
    ws = openpyxl.load_workbook(path).active
    return [tuple(c.value for c in fila) for fila in ws.iter_rows()]


def test_exportar_setmana_capçalera_iso_i_total(
    nedador_test, mesocicle_test, estructura_test, tmp_path
):
    """Una setmana: pestanya pròpia, capçalera amb setmana ISO 41 i fila Total."""
    microcicle = mesocicle_test.microcicles[1]  # 05-11/10/2026
    sessions = [
        _sessio(2, "dimarts", estructura_test),
        _sessio(2, "dilluns", estructura_test),
    ]

    path = exportar_setmana_excel(
        nedador_test, mesocicle_test, microcicle, sessions, tmp_path / "setmana.xlsx"
    )

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert ws.title == "Setmana 2"
    assert "Setmana 41" in ws.cell(row=1, column=1).value
    assert "05-11/10/2026" in ws.cell(row=1, column=1).value

    valors = _valors(path)
    dies = [f[0] for f in valors if f[0] in ("Dilluns 5", "Dimarts 6")]
    assert dies == ["Dilluns 5", "Dimarts 6"]  # ordenats dilluns..diumenge
    totals = [f[8] for f in valors if f[1] == "Total"]
    assert totals == [400, 400]


def test_exportar_setmana_igual_que_dins_del_mesocicle(
    nedador_test, mesocicle_test, estructura_test, tmp_path
):
    """El full d'una setmana té les mateixes files que aquella setmana sola al mesocicle."""
    microcicle = mesocicle_test.microcicles[1]
    sessions = [_sessio(2, "dilluns", estructura_test)]

    path_setmana = exportar_setmana_excel(
        nedador_test, mesocicle_test, microcicle, sessions, tmp_path / "s.xlsx"
    )
    path_meso = exportar_mesocicle_excel(
        nedador_test, mesocicle_test, {2: sessions}, tmp_path / "m.xlsx"
    )

    assert _valors(path_setmana) == _valors(path_meso)


# --- Fase H: rol a la capçalera i rutina d'espatlla ---


def test_exportar_setmana_rol_i_rutina_espatlla(mesocicle_test, estructura_test, tmp_path):
    """Sessions de plantilla: etiqueta del rol; rutina d'espatlla el dimecres (descans)."""
    nedador = Nedador(
        id="jep", nom="Jep", categoria="master", proves_objectiu=["100m lliure"],
        mode_ritme="temps", dies_disponibles=["dilluns", "dimarts", "dijous", "divendres"],
        rutina_espatlla_dia="dimecres",
    )
    microcicle = mesocicle_test.microcicles[1]  # 05-11/10/2026
    sessions = []
    for dia, rol in (("dilluns", "aerobica"), ("dimarts", "qualitat"), ("dijous", "tecnica")):
        s = _sessio(2, dia, estructura_test)
        s.rol = rol
        s.notes = "Setmana sense competició"
        sessions.append(s)

    path = exportar_setmana_excel(
        nedador, mesocicle_test, microcicle, sessions, tmp_path / "h.xlsx"
    )
    primeres = [f[0] for f in _valors(path) if f[0]]

    assert "Dilluns 5 — Aeròbic i tècnica" in primeres
    assert "Dimarts 6 — Qualitat" in primeres
    rutina = next(i for i, v in enumerate(primeres) if v.startswith("Dimecres 7 — Descans"))
    assert primeres.index("Dimarts 6 — Qualitat") < rutina < primeres.index("Dijous 8 — Tècnica")


def test_rutina_espatlla_no_el_dia_de_competicio(mesocicle_test, estructura_test, tmp_path):
    """Rutina el dissabte, però aquest dissabte hi ha competició: no surt."""
    nedador = Nedador(
        id="jep", nom="Jep", categoria="master", proves_objectiu=["100m lliure"],
        mode_ritme="temps", rutina_espatlla_dia="dissabte",
    )
    microcicle = mesocicle_test.microcicles[1].model_copy(update={"dia_competicio": "dissabte"})
    s = _sessio(2, "dilluns", estructura_test)
    s.rol, s.notes = "aerobica", "Competició el dissabte"

    path = exportar_setmana_excel(nedador, mesocicle_test, microcicle, [s], tmp_path / "c.xlsx")
    assert not any("Rutina d'espatlla" in str(f[0]) for f in _valors(path))

    microcicle_lliure = mesocicle_test.microcicles[1]
    path2 = exportar_setmana_excel(
        nedador, mesocicle_test, microcicle_lliure, [s], tmp_path / "d.xlsx"
    )
    assert any(str(f[0]).startswith("Dissabte 10 — Descans") for f in _valors(path2))
